"""Make a safe import preview from the Excel sources; never writes a database.

The preview contains sheet/row references and flags, but no customer names, full
phone numbers, CCCD or notes. Person groups are candidates matched strictly on
normalized name AND at least one normalized phone; applications are never merged.

Usage: python phase1/dry_run_import.py [--root PATH] [--out PATH]
"""

from __future__ import annotations

import argparse
import csv
import json
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

import openpyxl

from audit_workbooks import DAILY_FILE, NE_FILE, SOURCES, fold, has_value, phones, source_rows, write_csv


PROGRAM_COLUMNS = {
    "ne2026": 4,
    "ne2025": 3,
    "ne2024": 3,
    "data_ctv": 3,
    "lead_self": 6,
    "zalo_oa": 3,
    "web_gmail": 3,
}
TUITION_COLUMNS = {"ne2026": 19, "ne2025": 16, "ne2024": 12}


def csv_rows(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as file:
        return list(csv.DictReader(file))


class DisjointSet:
    def __init__(self, size: int) -> None:
        self.parent = list(range(size))
        self.rank = [0] * size

    def find(self, value: int) -> int:
        while self.parent[value] != value:
            self.parent[value] = self.parent[self.parent[value]]
            value = self.parent[value]
        return value

    def union(self, left: int, right: int) -> None:
        a, b = self.find(left), self.find(right)
        if a == b:
            return
        if self.rank[a] < self.rank[b]:
            a, b = b, a
        self.parent[b] = a
        if self.rank[a] == self.rank[b]:
            self.rank[a] += 1


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parent.parent)
    parser.add_argument("--out", type=Path, default=Path(__file__).resolve().parent / "reports")
    args = parser.parse_args()
    root, out = args.root.resolve(), args.out.resolve()
    out.mkdir(parents=True, exist_ok=True)

    sales_alias = {
        fold(row["source_name"]): row
        for row in csv_rows(out / "sales_alias_review.csv")
    }
    ctv_alias = {
        fold(row["source_ctv_name"]): row
        for row in csv_rows(out / "ctv_alias_review.csv")
    }
    books = {
        name: openpyxl.load_workbook(root / name, data_only=True, read_only=False)
        for name in (NE_FILE, DAILY_FILE)
    }

    records: list[dict[str, Any]] = []
    for spec in SOURCES:
        if spec.key == "ctv_roster":
            continue
        sheet = books[spec.workbook][spec.sheet]
        for row in source_rows(sheet, spec):
            owner = sheet.cell(row, spec.owner_col).value if spec.owner_col else None
            ctv = sheet.cell(row, spec.ctv_col).value if spec.ctv_col else None
            program_col = PROGRAM_COLUMNS.get(spec.key)
            program = sheet.cell(row, program_col).value if program_col else None
            normalized_phones = phones(sheet.cell(row, spec.phone_col).value) if spec.phone_col else set()
            amount_col = TUITION_COLUMNS.get(spec.key)
            amount = sheet.cell(row, amount_col).value if amount_col else None
            records.append(
                {
                    "source": spec.key,
                    "sheet": spec.sheet,
                    "row": row,
                    "entity": "NE_HISTORY" if spec.key.startswith("ne") else "LEAD",
                    "name_key": fold(sheet.cell(row, spec.name_col).value),
                    "phones": normalized_phones,
                    "alternative_contact_present": spec.key == "lead_online" and has_value(sheet.cell(row, 4).value),
                    "program_present": has_value(program),
                    "owner_raw_present": has_value(owner),
                    "owner_expected": spec.owner_col is not None,
                    "owner_alias": sales_alias.get(fold(owner)) if has_value(owner) else None,
                    "ctv_raw_present": has_value(ctv),
                    "ctv_alias": ctv_alias.get(fold(ctv)) if has_value(ctv) else None,
                    "tuition_positive_numeric": isinstance(amount, (int, float)) and not isinstance(amount, bool) and amount > 0,
                }
            )

    # Strict person candidate match: same normalized name and a shared phone.
    # Different names with a shared phone remain separate for manual review.
    sets = DisjointSet(len(records))
    seen: dict[tuple[str, str], int] = {}
    for idx, record in enumerate(records):
        for phone in record["phones"]:
            key = (record["name_key"], phone)
            if key in seen:
                sets.union(idx, seen[key])
            else:
                seen[key] = idx
    roots = sorted({sets.find(idx) for idx in range(len(records))})
    group_id = {root: f"G{position:05d}" for position, root in enumerate(roots, 1)}
    group_members: dict[str, list[int]] = defaultdict(list)
    for idx in range(len(records)):
        group_members[group_id[sets.find(idx)]].append(idx)

    preview: list[dict[str, Any]] = []
    for idx, record in enumerate(records):
        flags: list[str] = []
        if not record["phones"]:
            flags.append("CONTACT_VIA_FB_ONLY" if record["alternative_contact_present"] else "NO_VALID_CONTACT")
        if record["entity"] == "NE_HISTORY" and not record["program_present"]:
            flags.append("MISSING_PROGRAM")
        if record["owner_expected"] and not record["owner_raw_present"]:
            flags.append("MISSING_OWNER")
        if record["entity"] == "NE_HISTORY" and not record["owner_expected"]:
            flags.append("HISTORICAL_OWNER_MAPPING_PENDING")
        if record["source"] == "data_ctv" and not record["ctv_raw_present"]:
            flags.append("MISSING_CTV")
        if record["owner_raw_present"] and record["owner_alias"] is None:
            flags.append("OWNER_ALIAS_NOT_IN_REVIEW")
        if record["ctv_raw_present"] and record["ctv_alias"] is None:
            flags.append("CTV_ALIAS_NOT_IN_REVIEW")

        owner_row = record["owner_alias"]
        ctv_row = record["ctv_alias"]
        owner_approved = bool(owner_row and owner_row.get("approved_user_id") and owner_row.get("decision") == "DUYET")
        ctv_approved = bool(ctv_row and ctv_row.get("approved_ctv_id") and ctv_row.get("decision") == "DUYET")
        if record["owner_raw_present"] and not owner_approved:
            flags.append("OWNER_APPROVAL_PENDING")
        if record["ctv_raw_present"] and not ctv_approved:
            flags.append("CTV_APPROVAL_PENDING")
        if record["entity"] == "NE_HISTORY":
            flags.append("PAYMENT_VERIFICATION_PENDING")

        core_issue = any(flag in flags for flag in (
            "NO_VALID_CONTACT", "MISSING_PROGRAM", "MISSING_OWNER", "MISSING_CTV",
            "OWNER_ALIAS_NOT_IN_REVIEW", "CTV_ALIAS_NOT_IN_REVIEW",
        ))
        decision = "REVIEW_STRUCTURE" if core_issue else "MAPPING_OR_PAYMENT_APPROVAL_PENDING" if flags else "STRUCTURE_READY"
        preview.append(
            {
                "source": record["source"],
                "sheet": record["sheet"],
                "row": record["row"],
                "person_candidate_group": group_id[sets.find(idx)],
                "person_group_rows": len(group_members[group_id[sets.find(idx)]]),
                "entity": record["entity"],
                "phone_count": len(record["phones"]),
                "alternative_contact_present": record["alternative_contact_present"],
                "program_present": record["program_present"],
                "owner_alias_suggestion": owner_row.get("suggested_group", "") if owner_row else "",
                "ctv_roster_name_suggestion": ctv_row.get("exact_roster_name_after_normalization", "") if ctv_row else "",
                "historical_tuition_positive_numeric_unverified": record["tuition_positive_numeric"] if record["entity"] == "NE_HISTORY" else "",
                "decision": decision,
                "flags": "|".join(flags),
            }
        )

    link_candidates: list[dict[str, Any]] = []
    for group, indices in group_members.items():
        kinds = {records[idx]["entity"] for idx in indices}
        if kinds != {"NE_HISTORY", "LEAD"}:
            continue
        ne_count = sum(records[idx]["entity"] == "NE_HISTORY" for idx in indices)
        lead_count = len(indices) - ne_count
        for idx in indices:
            record = records[idx]
            link_candidates.append(
                {
                    "person_candidate_group": group,
                    "source": record["source"],
                    "sheet": record["sheet"],
                    "row": record["row"],
                    "entity": record["entity"],
                    "ne_rows_in_group": ne_count,
                    "lead_rows_in_group": lead_count,
                    "application_link_decision": "CHUA_DUYET",
                }
            )

    write_csv(
        out / "import_preview.csv",
        [
            "source", "sheet", "row", "person_candidate_group", "person_group_rows", "entity",
            "phone_count", "alternative_contact_present", "program_present", "owner_alias_suggestion", "ctv_roster_name_suggestion",
            "historical_tuition_positive_numeric_unverified", "decision", "flags",
        ],
        preview,
    )
    write_csv(
        out / "lead_ne_link_candidates.csv",
        [
            "person_candidate_group", "source", "sheet", "row", "entity",
            "ne_rows_in_group", "lead_rows_in_group", "application_link_decision",
        ],
        link_candidates,
    )
    summary = {
        "source_rows_previewed": len(records),
        "strict_person_candidate_groups": len(group_members),
        "groups_with_multiple_source_rows": sum(len(indices) > 1 for indices in group_members.values()),
        "lead_ne_link_candidate_groups": len({row["person_candidate_group"] for row in link_candidates}),
        "lead_ne_link_candidate_rows": len(link_candidates),
        "decisions": dict(Counter(row["decision"] for row in preview)),
        "flag_counts": dict(Counter(flag for row in preview for flag in row["flags"].split("|") if flag)),
        "note": "Đây là preview trong bộ nhớ, không tạo database, không công nhận NE và không ghép application tự động.",
    }
    (out / "import_dry_run_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False))


if __name__ == "__main__":
    main()
