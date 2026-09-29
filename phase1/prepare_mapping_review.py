"""Prepare staff/CTV alias review sheets. Includes business contact names, no students or phones.

Usage: python phase1/prepare_mapping_review.py [--root PATH] [--out PATH]
The CSVs are for human approval; suggestions must not be imported as approved IDs.
"""

from __future__ import annotations

import argparse
import csv
import re
import unicodedata
from collections import Counter
from pathlib import Path

import openpyxl

from audit_workbooks import DAILY_FILE, NE_FILE, SOURCES, source_rows


def fold(value: str) -> str:
    value = unicodedata.normalize("NFD", value.casefold().replace("đ", "d"))
    value = "".join(ch for ch in value if unicodedata.category(ch) != "Mn")
    return re.sub(r"[^a-z0-9]+", " ", value).strip()


def write_csv(path: Path, headers: list[str], rows: list[dict[str, object]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)


def previous_by_key(path: Path, key: str) -> dict[str, dict[str, str]]:
    if not path.exists():
        return {}
    with path.open(encoding="utf-8-sig", newline="") as file:
        return {row[key]: row for row in csv.DictReader(file)}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parent.parent)
    parser.add_argument("--out", type=Path, default=Path(__file__).resolve().parent / "reports")
    args = parser.parse_args()
    out = args.out.resolve()
    out.mkdir(parents=True, exist_ok=True)

    ne_book = openpyxl.load_workbook(args.root / NE_FILE, data_only=True)
    ne = ne_book["NE 2026"]
    ne_owner_counts = Counter(
        str(ne.cell(row, 28).value).strip()
        for row in range(4, ne.max_row + 1)
        if ne.cell(row, 3).value is not None and ne.cell(row, 28).value is not None
    )
    daily = openpyxl.load_workbook(args.root / DAILY_FILE, data_only=True)
    books = {NE_FILE: ne_book, DAILY_FILE: daily}
    owner_counts = Counter()
    owner_sources: dict[str, set[str]] = {}
    for spec in SOURCES:
        if spec.owner_col is None:
            continue
        sheet = books[spec.workbook][spec.sheet]
        for row in source_rows(sheet, spec):
            value = sheet.cell(row, spec.owner_col).value
            if value is None or str(value).strip() == "":
                continue
            name = str(value).strip()
            owner_counts[name] += 1
            owner_sources.setdefault(name, set()).add(spec.key)
    suggestions = {
        "DANG ANH THU": "Thư",
        "THU": "Thư",
        "THUY DUONG": "Dương",
        "HOANG THI THUY DUONG": "Dương",
        "DUONG": "Dương",
        "NHU BINH": "Bình",
        "TRAN THI NHU BINH": "Bình",
        "BINH": "Bình",
        "PHUONG TRINH": "Trinh",
        "TRINH": "Trinh",
        "NGUYEN THI MINH TINH": "Tính",
        "MINH TINH": "Tính",
        "TINH": "Tính",
        "THUY NGA": "Nga",
        "NGA": "Nga",
    }
    old_owner = previous_by_key(out / "sales_alias_review.csv", "source_name")
    owner_rows = [
        {
            "source_name": name,
            "all_source_rows": count,
            "ne2026_rows": ne_owner_counts[name],
            "sources": " | ".join(sorted(owner_sources.get(name, set()))),
            "suggested_group": suggestions.get(fold(name).upper(), ""),
            "approved_user_id": old_owner.get(name, {}).get("approved_user_id", ""),
            "decision": old_owner.get(name, {}).get("decision", "CHƯA_DUYỆT"),
            "review_note": old_owner.get(name, {}).get("review_note", ""),
        }
        for name, count in owner_counts.most_common()
    ]
    write_csv(
        out / "sales_alias_review.csv",
        ["source_name", "all_source_rows", "ne2026_rows", "sources", "suggested_group", "approved_user_id", "decision", "review_note"],
        owner_rows,
    )

    roster = daily["DANH SÁCH CTV "]
    roster_names = [
        str(roster.cell(row, 2).value).strip()
        for row in range(3, roster.max_row + 1)
        if roster.cell(row, 2).value is not None
    ]
    roster_by_key: dict[str, list[str]] = {}
    for name in roster_names:
        roster_by_key.setdefault(fold(name), []).append(name)
    data = daily["Data CTV"]
    ctv_counts = Counter(
        str(data.cell(row, 9).value).strip()
        for row in range(2, data.max_row + 1)
        if data.cell(row, 2).value is not None and data.cell(row, 9).value is not None
    )
    tele_by_ctv: dict[str, Counter[str]] = {}
    for row in range(2, data.max_row + 1):
        ctv_name = data.cell(row, 9).value
        lead_name = data.cell(row, 2).value
        tele_name = data.cell(row, 7).value
        if lead_name is None or ctv_name is None or tele_name is None:
            continue
        tele_by_ctv.setdefault(str(ctv_name).strip(), Counter())[str(tele_name).strip()] += 1

    def tele_hint(name: str) -> tuple[str, int, int]:
        counts = tele_by_ctv.get(name, Counter())
        if not counts:
            return "", 0, 0
        top, count = counts.most_common(1)[0]
        return top, count, sum(counts.values())

    old_ctv = previous_by_key(out / "ctv_alias_review.csv", "source_ctv_name")
    ctv_rows = [
        {
            "source_ctv_name": name,
            "data_ctv_rows": count,
            "exact_roster_name_after_normalization": " | ".join(roster_by_key.get(fold(name), [])),
            "suggested_sales_from_tele": tele_hint(name)[0],
            "supporting_leads": tele_hint(name)[1],
            "leads_with_tele": tele_hint(name)[2],
            "approved_ctv_id": old_ctv.get(name, {}).get("approved_ctv_id", ""),
            "owner_sales_user_id": old_ctv.get(name, {}).get("owner_sales_user_id", ""),
            "decision": old_ctv.get(name, {}).get("decision", "CHƯA_DUYỆT"),
            "review_note": old_ctv.get(name, {}).get("review_note", ""),
        }
        for name, count in ctv_counts.most_common()
    ]
    old_roster = previous_by_key(out / "ctv_roster_review.csv", "roster_name")
    write_csv(
        out / "ctv_alias_review.csv",
        [
            "source_ctv_name", "data_ctv_rows", "exact_roster_name_after_normalization",
            "suggested_sales_from_tele", "supporting_leads", "leads_with_tele",
            "approved_ctv_id", "owner_sales_user_id", "decision", "review_note",
        ],
        ctv_rows,
    )
    write_csv(
        out / "ctv_roster_review.csv",
        ["roster_name", "approved_ctv_id", "owner_sales_user_id", "status", "review_note"],
        [
            {
                "roster_name": name,
                "approved_ctv_id": old_roster.get(name, {}).get("approved_ctv_id", ""),
                "owner_sales_user_id": old_roster.get(name, {}).get("owner_sales_user_id", ""),
                "status": old_roster.get(name, {}).get("status", "CHƯA_DUYỆT"),
                "review_note": old_roster.get(name, {}).get("review_note", ""),
            }
            for name in roster_names
        ],
    )
    kpi = daily["Report NE 2026"]
    sales_names = [str(kpi.cell(3, col).value).strip() for col in range(2, 7) if kpi.cell(3, col).value is not None]
    old_sales_roster = previous_by_key(out / "sales_roster_draft.csv", "source_name")
    write_csv(
        out / "sales_roster_draft.csv",
        ["source_name", "source_sheet", "proposed_status", "approved_user_id", "decision", "review_note"],
        [
            {
                "source_name": name,
                "source_sheet": "Report NE 2026, row 3",
                "proposed_status": "Trong KPI 2026",
                "approved_user_id": old_sales_roster.get(name, {}).get("approved_user_id", ""),
                "decision": old_sales_roster.get(name, {}).get("decision", "CHƯA_DUYỆT"),
                "review_note": old_sales_roster.get(name, {}).get("review_note", ""),
            }
            for name in sales_names
        ],
    )
    print(f"Created {len(owner_rows)} sales aliases, {len(sales_names)} KPI sales, {len(ctv_rows)} CTV aliases, {len(roster_names)} roster rows in {out}")


if __name__ == "__main__":
    main()
