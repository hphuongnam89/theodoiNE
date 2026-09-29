"""Read-only audit of the two source workbooks. Outputs contain no full names or phones.

Usage: python phase1/audit_workbooks.py [--root PATH] [--out PATH]
Requires openpyxl. Never writes to the source XLSX files.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl


NE_FILE = "1. THEO DÕI NE.xlsx"
DAILY_FILE = "4. Report Daily.xlsx"


@dataclass(frozen=True)
class Source:
    key: str
    workbook: str
    sheet: str
    first_row: int
    name_col: int
    phone_col: int | None
    status_col: int | None = None
    owner_col: int | None = None
    ctv_col: int | None = None


SOURCES = (
    Source("ne2026", NE_FILE, "NE 2026", 4, 3, 8, owner_col=28),
    Source("ne2025", NE_FILE, "NE 2025", 4, 2, 7),
    Source("ne2024", NE_FILE, "NE 2024 ", 4, 2, None),
    Source("data_ctv", DAILY_FILE, "Data CTV", 2, 2, 4, 6, 7, 9),
    Source("lead_online", DAILY_FILE, "LEAD ONLINE", 2, 2, 3, 7, 6),
    Source("lead_self", DAILY_FILE, "Lead tự tạo", 2, 4, 7, 9, 10),
    Source("zalo_oa", DAILY_FILE, "Zalo OA", 2, 2, 4, 6, 7),
    Source("web_gmail", DAILY_FILE, "Web gmail Viện", 2, 2, 4, 6, 7),
    Source("ctv_roster", DAILY_FILE, "DANH SÁCH CTV ", 3, 2, 5),
)


def has_value(value: Any) -> bool:
    return value is not None and str(value).strip() != ""


def fold(value: Any) -> str:
    text = str(value or "").casefold().replace("đ", "d")
    text = unicodedata.normalize("NFD", text)
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


PHONE_PATTERN = re.compile(r"(?<!\d)(?:\+?84|0)?(?:[\s.()-]*\d){9,10}(?!\d)")


def phones(value: Any) -> set[str]:
    if not has_value(value):
        return set()
    raw = str(int(value)) if isinstance(value, float) and value.is_integer() else str(value)
    result: set[str] = set()
    for token in PHONE_PATTERN.findall(raw):
        number = re.sub(r"\D", "", token)
        if number.startswith("84") and len(number) in (11, 12):
            number = "0" + number[2:]
        elif len(number) == 9:
            number = "0" + number
        if len(number) in (10, 11) and number.startswith("0") and len(set(number)) > 1:
            result.add(number)
    return result


def write_csv(path: Path, fields: list[str], rows: list[dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def source_rows(sheet: Any, spec: Source) -> list[int]:
    return [
        row
        for row in range(spec.first_row, sheet.max_row + 1)
        if has_value(sheet.cell(row, spec.name_col).value)
    ]


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parent.parent)
    parser.add_argument("--out", type=Path, default=Path(__file__).resolve().parent / "reports")
    args = parser.parse_args()
    root, out = args.root.resolve(), args.out.resolve()
    out.mkdir(parents=True, exist_ok=True)

    for name in (NE_FILE, DAILY_FILE):
        if not (root / name).is_file():
            raise FileNotFoundError(root / name)

    books = {
        name: openpyxl.load_workbook(root / name, data_only=True, read_only=False)
        for name in (NE_FILE, DAILY_FILE)
    }
    formula_book = openpyxl.load_workbook(root / DAILY_FILE, data_only=False, read_only=False)

    profile: dict[str, Any] = {
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "source_files": {
            name: {
                "size_bytes": (root / name).stat().st_size,
                "sha256": hashlib.sha256((root / name).read_bytes()).hexdigest(),
            }
            for name in (NE_FILE, DAILY_FILE)
        },
        "sources": {}, "ne2026": {}, "kpi_2026": {}, "ctv": {},
    }
    review: list[dict[str, Any]] = []
    phone_index: dict[str, list[dict[str, str]]] = defaultdict(list)
    roster_names: set[str] = set()

    for spec in SOURCES:
        sheet = books[spec.workbook][spec.sheet]
        rows = source_rows(sheet, spec)
        status = Counter()
        phone_rows = 0
        owner_present = 0
        for row in rows:
            ref = f"{spec.sheet}!{row}"
            if spec.status_col:
                value = sheet.cell(row, spec.status_col).value
                status[str(value).strip() if has_value(value) else "<trống>"] += 1
            if spec.owner_col and has_value(sheet.cell(row, spec.owner_col).value):
                owner_present += 1
            elif spec.owner_col:
                review.append({"issue": "MISSING_SALES_OWNER", "source": spec.key, "sheet": spec.sheet, "row": row, "detail": ""})
            if spec.phone_col:
                found = phones(sheet.cell(row, spec.phone_col).value)
                phone_rows += bool(found)
                for number in found:
                    phone_index[number].append(
                        {"source": spec.key, "sheet": spec.sheet, "row": str(row), "name_key": fold(sheet.cell(row, spec.name_col).value)}
                    )
            if spec.key == "ctv_roster":
                roster_names.add(fold(sheet.cell(row, spec.name_col).value))
            if spec.key == "ne2026":
                if not has_value(sheet.cell(row, 4).value):
                    review.append({"issue": "MISSING_PROGRAM", "source": spec.key, "sheet": spec.sheet, "row": row, "detail": ""})
                amount = sheet.cell(row, 19).value
                if not (isinstance(amount, (int, float)) and not isinstance(amount, bool) and amount > 0):
                    review.append({"issue": "TUITION_NOT_POSITIVE_NUMBER", "source": spec.key, "sheet": spec.sheet, "row": row, "detail": "Giá trị gốc cần xem tại Excel"})

        profile["sources"][spec.key] = {
            "sheet": spec.sheet,
            "rows_with_name": len(rows),
            "rows_with_valid_phone": phone_rows,
            "rows_with_owner": owner_present if spec.owner_col else None,
            "status_counts": dict(status.most_common()) if spec.status_col else None,
        }

    ne = books[NE_FILE]["NE 2026"]
    ne_rows = source_rows(ne, SOURCES[0])
    tuition = [ne.cell(row, 19).value for row in ne_rows]
    profile["ne2026"] = {
        "rows_with_name": len(ne_rows),
        "rows_with_program": sum(has_value(ne.cell(row, 4).value) for row in ne_rows),
        "rows_with_sales_owner_in_AB": sum(has_value(ne.cell(row, 28).value) for row in ne_rows),
        "tuition_cells_nonempty": sum(has_value(value) for value in tuition),
        "tuition_cells_positive_numeric_unverified": sum(
            isinstance(value, (int, float)) and not isinstance(value, bool) and value > 0 for value in tuition
        ),
        "summary_2026_cached_count": books[NE_FILE]["SUMMARY 2026"].cell(2, 3).value,
    }

    ctv_sheet = books[DAILY_FILE]["Data CTV"]
    ctv_rows = source_rows(ctv_sheet, next(source for source in SOURCES if source.key == "data_ctv"))
    ctv_filled = 0
    ctv_matched = 0
    ctv_values: set[str] = set()
    for row in ctv_rows:
        value = ctv_sheet.cell(row, 9).value
        if not has_value(value):
            review.append({"issue": "MISSING_CTV", "source": "data_ctv", "sheet": "Data CTV", "row": row, "detail": ""})
            continue
        ctv_filled += 1
        key = fold(value)
        ctv_values.add(key)
        if key in roster_names:
            ctv_matched += 1
        else:
            review.append({"issue": "CTV_NAME_NOT_IN_ROSTER", "source": "data_ctv", "sheet": "Data CTV", "row": row, "detail": "Cần đối chiếu tên viết tắt/mã CTV"})
    profile["ctv"] = {
        "data_ctv_rows": len(ctv_rows),
        "data_ctv_rows_with_ctv": ctv_filled,
        "distinct_ctv_names_normalized": len(ctv_values),
        "roster_names": len(roster_names),
        "rows_exact_name_match_after_normalization": ctv_matched,
    }

    duplicate_rows: list[dict[str, Any]] = []
    groups = []
    for number, refs in phone_index.items():
        unique = {(ref["source"], ref["row"]): ref for ref in refs}
        if len(unique) >= 2:
            groups.append((number, sorted(unique.values(), key=lambda x: (x["source"], int(x["row"])))))
    groups.sort(key=lambda item: (item[1][0]["source"], int(item[1][0]["row"]), item[0]))
    cross_source = 0
    for idx, (number, refs) in enumerate(groups, 1):
        sources = {ref["source"] for ref in refs}
        cross_source += len(sources) > 1
        same_name = len({ref["name_key"] for ref in refs}) == 1
        for ref in refs:
            duplicate_rows.append(
                {
                    "candidate_group": f"P{idx:04d}",
                    "masked_phone": "******" + number[-4:],
                    "source": ref["source"],
                    "sheet": ref["sheet"],
                    "row": ref["row"],
                    "group_size": len(refs),
                    "cross_source": len(sources) > 1,
                    "same_normalized_name": same_name,
                }
            )
    profile["phone_duplicate_candidates"] = {
        "groups": len(groups),
        "groups_cross_source": cross_source,
        "rows_in_groups": len(duplicate_rows),
        "note": "Số điện thoại trùng là ứng viên cần kiểm tra, không tự động ghép bản ghi.",
    }

    kpi = books[DAILY_FILE]["Report NE 2026"]
    kpi_formulas = formula_book["Report NE 2026"]
    kpi_rows: list[dict[str, Any]] = []
    for row in range(4, 16):
        month = kpi.cell(row, 1).value
        if not isinstance(month, (int, float)) or not 1 <= month <= 9:
            continue

        def numeric(value: Any) -> float:
            return float(value) if isinstance(value, (int, float)) and not isinstance(value, bool) else 0.0

        target = numeric(kpi.cell(row, 7).value)
        total = numeric(kpi.cell(row, 13).value)
        staff = sum(numeric(kpi.cell(row, col).value) for col in range(8, 13))
        kpi_rows.append(
            {
                "month": int(month),
                "target_G": target,
                "reported_total_M": total,
                "sum_staff_H_to_L": staff,
                "difference_M_minus_staff": total - staff,
                "formula_M": kpi_formulas.cell(row, 13).value,
            }
        )
    profile["kpi_2026"] = {
        "months": "1-9",
        "target_total": sum(row["target_G"] for row in kpi_rows),
        "reported_total_M": sum(row["reported_total_M"] for row in kpi_rows),
        "sum_staff_H_to_L": sum(row["sum_staff_H_to_L"] for row in kpi_rows),
        "note": "Các công thức cột M tham chiếu phần báo cáo ngày phía dưới; một số dòng nhân sự bị bỏ qua.",
    }

    write_csv(out / "review_queue.csv", ["issue", "source", "sheet", "row", "detail"], review)
    write_csv(
        out / "duplicate_candidates.csv",
        ["candidate_group", "masked_phone", "source", "sheet", "row", "group_size", "cross_source", "same_normalized_name"],
        duplicate_rows,
    )
    write_csv(
        out / "kpi_reconciliation.csv",
        ["month", "target_G", "reported_total_M", "sum_staff_H_to_L", "difference_M_minus_staff", "formula_M"],
        kpi_rows,
    )
    (out / "source_profile.json").write_text(json.dumps(profile, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"output": str(out), "profiles": len(SOURCES), "review_items": len(review), **profile["phone_duplicate_candidates"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
